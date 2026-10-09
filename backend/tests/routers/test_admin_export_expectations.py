"""명시적인 정·역거래 쌍의 F704 기간별 실제 창고 증감 계약."""

from datetime import datetime, timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.models import TransactionLog, TransactionTypeEnum
from app.services.f704_02_ledger import collect_entries


def test_repeating_completed_integrity_repair_preserves_all_rows_and_single_audit(db_session, make_item):
    from sqlalchemy import select
    from app.database import Base
    from app.models import AdminAuditLog, ShippingRequestStatusEnum
    from app.services.inventory_integrity_repair import InventoryIntegrityRepairError, repair_inventory_integrity_issue
    from tests.services.test_inventory_integrity_repair import _workflow_residue

    request, issue = _workflow_residue(db_session, make_item)
    repaired = repair_inventory_integrity_issue(db_session, problem_id=issue.problem_id, approved_by="검수 관리자", apply=True)
    db_session.commit()
    assert repaired.applied is True
    assert request.status == ShippingRequestStatusEnum.CANCELLED
    assert db_session.query(AdminAuditLog).count() == 1
    before = {
        table.name: list(db_session.execute(select(table).order_by(*table.primary_key.columns)).mappings())
        for table in Base.metadata.sorted_tables
    }
    with pytest.raises(InventoryIntegrityRepairError, match="현재 진단 결과에서 문제 ID를 찾을 수 없습니다"):
        repair_inventory_integrity_issue(db_session, problem_id=issue.problem_id, approved_by="다시 요청한 관리자", apply=True)
    db_session.rollback()
    after = {
        table.name: list(db_session.execute(select(table).order_by(*table.primary_key.columns)).mappings())
        for table in Base.metadata.sorted_tables
    }
    assert after == before
    assert db_session.query(AdminAuditLog).count() == 1


@pytest.mark.parametrize(
    ("original_at", "reversal_at", "year", "expected_indices"),
    [
        # 같은 연도의 월경계도 각각의 KST 발생일로 유지한다.
        (datetime(2026, 9, 30, 14, 59), datetime(2026, 9, 30, 15), 2026, [0, 1]),
        # KST 자정: 원거래는 2026-12-31 23:59, 역거래는 2027-01-01 00:00.
        (datetime(2026, 12, 31, 14, 59), datetime(2026, 12, 31, 15), 2026, [0]),
        (datetime(2026, 12, 31, 14, 59), datetime(2026, 12, 31, 15), 2027, [1]),
    ],
)
def test_f704_explicit_reversal_pair_preserves_each_actual_warehouse_change_in_its_kst_year(
    client, db_session, make_item, original_at, reversal_at, year, expected_indices
):
    item = make_item(name="F704 explicit original and reversal")
    original = TransactionLog(
        item_id=item.item_id,
        transaction_type=TransactionTypeEnum.RECEIVE,
        quantity_change=4,
        quantity_before=0,
        quantity_after=4,
        warehouse_qty_before=0,
        warehouse_qty_after=4,
        inventory_effect=[{"scope": "warehouse", "delta": 4}],
        cancelled=True,
        cancelled_at=reversal_at,
        created_at=original_at,
    )
    db_session.add(original)
    db_session.flush()
    reversal = TransactionLog(
        item_id=item.item_id,
        transaction_type=TransactionTypeEnum.RECEIVE,
        quantity_change=-4,
        quantity_before=4,
        quantity_after=0,
        warehouse_qty_before=4,
        warehouse_qty_after=0,
        inventory_effect=[{"scope": "warehouse", "delta": -4}],
        reverses_log_id=original.log_id,
        created_at=reversal_at,
    )
    db_session.add(reversal)
    db_session.commit()

    pair = [original, reversal]
    selected = [pair[index] for index in expected_indices]
    before = [(row.cancelled, row.quantity_change, row.inventory_effect) for row in pair]
    entries = collect_entries(db_session, year)
    assert [row.log_id for row in entries] == [str(row.log_id) for row in selected]
    assert [row.quantity for row in entries] == [4 for _ in selected]
    assert [row.direction for row in entries] == ["입고" if row.quantity_change > 0 else "출고" for row in selected]

    response = client.get(
        f"/api/admin/audit-ledger/f704-02.xlsx?year={year}",
        headers={"X-Admin-Pin": "0000"},
    )
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content), data_only=True)
    try:
        rows = [row for row in book["양식"].iter_rows(min_row=4, values_only=True) if row[4] == item.item_name]
        assert len(rows) == len(selected)
        assert [row[5] for row in rows] == [4 for _ in selected]
        assert [row[6] for row in rows] == ["입고" if row.quantity_change > 0 else "출고" for row in selected]
        assert [row[1].date() for row in rows] == [(row.created_at + timedelta(hours=9)).date() for row in selected]
    finally:
        book.close()
    db_session.expire_all()
    assert [(row.cancelled, row.quantity_change, row.inventory_effect) for row in pair] == before


def test_f704_every_recorded_warehouse_effect_type_is_included_once_and_other_cells_are_excluded(
    db_session, make_item
):
    item = make_item(name="F704 all actual warehouse effect types")
    warehouse_logs = []
    for index, kind in enumerate(TransactionTypeEnum):
        delta = 1 if index % 2 == 0 else -1
        log = TransactionLog(
            item_id=item.item_id, transaction_type=kind, quantity_change=delta,
            inventory_effect=[{"scope": "warehouse", "delta": delta}],
            created_at=datetime(2026, 10, 1, 0, index),
        )
        db_session.add(log)
        warehouse_logs.append(log)
    db_session.add(TransactionLog(
        item_id=item.item_id, transaction_type=TransactionTypeEnum.RECEIVE,
        quantity_change=5,
        inventory_effect=[{"scope": "location", "department": "조립", "status": "PRODUCTION", "delta": 5}],
        created_at=datetime(2026, 10, 1, 1),
    ))
    db_session.commit()
    entries = collect_entries(db_session, 2026)
    assert [entry.log_id for entry in entries] == [str(log.log_id) for log in warehouse_logs]
    assert len({entry.log_id for entry in entries}) == len(list(TransactionTypeEnum))
    assert [entry.quantity for entry in entries] == [1 for _ in warehouse_logs]
    assert [entry.direction for entry in entries] == ["입고" if log.quantity_change > 0 else "출고" for log in warehouse_logs]
