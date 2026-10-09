"""Shipping PF inclusion in the verified weekly path, without frozen code edits."""
from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import event

from app.models import InventoryOperationRoleEnum, ProductSymbol, ShippingRequest, TransactionTypeEnum
from app.services.f705_02_production_log import collect_daily_quantities, render_workbook
from tests.routers.test_weekly_report import (
    WEEK_END,
    WEEK_START,
    _activate_verified_weekly_report,
    _add_operation_log,
    _add_snapshot,
    _dec,
    _make_prod_item,
)


def test_verified_pf_matrix_excludes_companion_unlinked_ship_and_pf_production(client, db_session) -> None:
    """Only the final PF pickup contributes, even when excluded flows share a model."""
    db_session.add(ProductSymbol(slot=1, symbol="3", model_name="DX3000"))
    db_session.flush()
    final_pf = _make_prod_item(db_session, name="Final PF", process_code="PF", model_symbol="3")
    companion = _make_prod_item(db_session, name="Companion PF", process_code="PF", model_symbol="3")
    request = ShippingRequest(base_pf_item_id=final_pf.item_id, final_pf_item_id=final_pf.item_id, request_quantity=5)
    db_session.add(request)
    db_session.flush()
    _activate_verified_weekly_report(db_session)
    for item, quantity, kind, linked in (
        (final_pf, 19, TransactionTypeEnum.PRODUCE, False),
        (companion, 11, TransactionTypeEnum.PRODUCE, False),
        (final_pf, -5, TransactionTypeEnum.SHIP, True),
        (companion, -2, TransactionTypeEnum.SHIP, True),
        (final_pf, -7, TransactionTypeEnum.SHIP, False),
    ):
        _add_operation_log(db_session, item=item, tx_type=kind,
            role=InventoryOperationRoleEnum.PRODUCT_OUTPUT if quantity > 0 else InventoryOperationRoleEnum.PRIMARY,
            quantity_change=quantity,
            effects=[{"scope": "location", "department": "출하", "status": "PRODUCTION", "delta": quantity}],
            action="produce" if quantity > 0 else "pickup", display_label="PF scope fixture",
            shipping_phase="PICKUP" if quantity < 0 else None,
            shipping_request_id=request.request_id if linked else None, domain="shipping" if quantity < 0 else "production")
    _add_snapshot(db_session, week_end=date(2026, 5, 3), item_quantities=[(final_pf, _dec(0)), (companion, _dec(0))], verified=True)
    _add_snapshot(db_session, week_end=date(2026, 5, 10), item_quantities=[(final_pf, _dec(7)), (companion, _dec(9))], verified=True)
    db_session.commit()
    response = client.get(f"/api/inventory/weekly-report?week_start={WEEK_START}&week_end={WEEK_END}")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["report_status"] == "verified"
    row = next(row for row in body["production_matrix"] if row["model_key"] == "DX3000")
    assert row["pf_qty"] == 5
    assert row["total_qty"] == 5


def test_verified_closed_week_pf_and_f705_keep_pickup_cancelled_next_week(client, db_session) -> None:
    """KST Sunday closure fixes the same PF quantity in the matrix and daily file."""
    db_session.add(ProductSymbol(slot=1, symbol="3", model_name="DX3000"))
    db_session.flush()
    final_pf = _make_prod_item(db_session, name="Closed-week PF", process_code="PF", model_symbol="3")
    request = ShippingRequest(base_pf_item_id=final_pf.item_id, final_pf_item_id=final_pf.item_id, request_quantity=5)
    db_session.add(request)
    db_session.flush()
    _activate_verified_weekly_report(db_session)
    pickup = _add_operation_log(db_session, item=final_pf, tx_type=TransactionTypeEnum.SHIP,
        role=InventoryOperationRoleEnum.PRIMARY, quantity_change=-5,
        effects=[{"scope": "location", "department": "출하", "status": "PRODUCTION", "delta": -5}],
        action="pickup", display_label="closed-week pickup", at=datetime(2026, 5, 8, 3),
        shipping_phase="PICKUP", shipping_request_id=request.request_id, domain="shipping")
    pickup.cancelled = True
    pickup.cancelled_at = datetime(2026, 5, 10, 15)  # Monday 00:00 KST: after closure.
    reversal = _add_operation_log(db_session, item=final_pf, tx_type=TransactionTypeEnum.SHIP,
        role=InventoryOperationRoleEnum.PRIMARY, quantity_change=5,
        effects=[{"scope": "location", "department": "출하", "status": "PRODUCTION", "delta": 5}],
        action="pickup-cancel", display_label="next-week reversal", at=pickup.cancelled_at,
        shipping_phase="PICKUP", shipping_request_id=request.request_id, domain="shipping")
    reversal.reverses_log_id = pickup.log_id
    _add_snapshot(db_session, week_end=date(2026, 5, 3), item_quantities=[(final_pf, _dec(5))], verified=True)
    _add_snapshot(db_session, week_end=date(2026, 5, 10), item_quantities=[(final_pf, _dec(0))], verified=True)
    db_session.commit()
    response = client.get(f"/api/inventory/weekly-report?week_start={WEEK_START}&week_end={WEEK_END}")
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["report_status"] == "verified"
    row = next(row for row in body["production_matrix"] if row["model_key"] == "DX3000")
    assert row["pf_qty"] == 5
    queries: list[str] = []
    def observe(_connection, _cursor, statement, _parameters, _context, _many) -> None:
        if "JOIN shipping_requests" in statement:
            queries.append(statement)
    engine = db_session.get_bind()
    event.listen(engine, "before_cursor_execute", observe)
    try:
        quantities = collect_daily_quantities(db_session, 2026)
    finally:
        event.remove(engine, "before_cursor_execute", observe)
    assert len(queries) == 1
    assert quantities[date(2026, 5, 8)][("PF", "DX3000")] == row["pf_qty"]
    workbook = load_workbook(BytesIO(render_workbook(2026, quantities)), data_only=True)
    assert workbook.worksheets[4].cell(27, 11).value == 5
