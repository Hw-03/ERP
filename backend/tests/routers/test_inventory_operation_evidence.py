"""Typed operation history includes observations and approved correction lineage."""

import csv
import json
from datetime import datetime, timedelta
from io import BytesIO, StringIO

from openpyxl import load_workbook

from app import models
from app.services import inventory_operations, stocktake
from tests.services.test_business_corrections import approval_preview, approve, requested, shipping_source
from tests.services.test_business_correction_boundaries import _reservation_only_prepare
from tests.services.test_stocktake_application import apply, confirm, preview
from tests.services.test_stocktake_counts import setup_round


def _range() -> dict[str, str]:
    """Select the current KST business day around service-created UTC timestamps."""
    today = (datetime.utcnow() + timedelta(hours=9)).date().isoformat()
    return {"start_date": today, "end_date": today}


def _csv_rows(client) -> list[dict[str, str]]:
    response = client.get("/api/inventory/operations/export.csv", params=_range())
    assert response.status_code == 200, response.text
    return list(csv.DictReader(StringIO(response.text)))


def _xlsx_rows(client) -> list[dict[str, object]]:
    response = client.get("/api/inventory/operations/export.xlsx", params=_range())
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    values = workbook.active.iter_rows(values_only=True)
    headings = next(values)
    return [dict(zip(headings, row)) for row in values]


def test_zero_difference_observation_is_visible_in_list_detail_and_exports(
    client, db_session, make_item,
) -> None:
    env = setup_round(db_session, make_item)
    confirmed = confirm(db_session, env)
    result = apply(db_session, env, preview(db_session, env, confirmed))
    item_id = str(env[2].item_id)
    operation_id = result["operation_id"]
    assert db_session.query(models.TransactionLog).count() == 0

    listing = client.get("/api/inventory/operations", params={"item_id": item_id})
    assert listing.status_code == 200, listing.text
    operation = next(row for row in listing.json()["items"] if row["operation_id"] == operation_id)
    assert operation["lines"] == operation["matching_lines"] == []
    observations = operation["matching_observation_lines"]
    assert len(observations) == 1
    assert observations[0]["row_type"] == "STOCKTAKE_OBSERVATION"
    assert observations[0]["item_id"] == item_id
    assert observations[0]["count_id"] == confirmed["count_id"]
    assert observations[0]["before_quantities"] == observations[0]["after_quantities"]
    assert observations[0]["observation_epoch_after"] == 1
    assert operation["observation_lines"] == observations

    detail = client.get(f"/api/inventory/operations/{operation_id}")
    assert detail.status_code == 200, detail.text
    assert detail.json()["observation_lines"] == observations
    for rows in (_csv_rows(client), _xlsx_rows(client)):
        exported = [row for row in rows if row["operation_id"] == operation_id]
        assert len(exported) == 1
        assert exported[0]["row_type"] == "STOCKTAKE_OBSERVATION"
        assert exported[0]["item_id"] == item_id
        assert exported[0]["source_id"] == result["application_id"]
        assert exported[0]["count_id"] == confirmed["count_id"]
        assert exported[0]["confirmation_revision_id"] == confirmed["confirmation_revision_id"]

    legacy = client.get("/api/inventory/transactions/export.csv", params=_range())
    assert legacy.status_code == 200, legacy.text
    assert list(csv.DictReader(StringIO(legacy.text))) == []


def test_operation_export_requires_bounded_kst_dates(client) -> None:
    for suffix in ("csv", "xlsx"):
        response = client.get(f"/api/inventory/operations/export.{suffix}")
        assert response.status_code == 400, response.text


def test_item_filter_keeps_only_matching_observation_lines(client, db_session, make_item) -> None:
    env = setup_round(db_session, make_item)
    other = make_item(name="Unrelated item", warehouse_qty=3)
    db_session.commit()
    confirmed = confirm(db_session, env)
    result = apply(db_session, env, preview(db_session, env, confirmed))
    listing = client.get("/api/inventory/operations", params={"item_id": str(other.item_id)})
    assert listing.status_code == 200, listing.text
    assert result["operation_id"] not in {row["operation_id"] for row in listing.json()["items"]}


def test_approved_business_correction_links_original_and_exports_without_log(
    client, db_session, make_item,
) -> None:
    env, _request, _allocation, source, _log = shipping_source(db_session, make_item)
    correction = requested(db_session, env, source)
    result = approve(db_session, env[1], correction, approval_preview(db_session, env[1], correction))
    operation_id = result["operation_id"]
    assert db_session.query(models.TransactionLog).count() == 1

    listing = client.get("/api/inventory/operations", params={"item_id": str(env[2].item_id)})
    assert listing.status_code == 200, listing.text
    operation = next(row for row in listing.json()["items"] if row["operation_id"] == operation_id)
    assert operation["business_correction"] == {
        "correction_id": correction["correction_id"],
        "source_type": "InventoryOperation",
        "source_id": str(source.operation_id),
        "source_operation_id": str(source.operation_id),
        "source_operation_ids": [str(source.operation_id)],
    }
    detail = client.get(f"/api/inventory/operations/{operation_id}")
    assert detail.status_code == 200, detail.text
    assert detail.json()["business_correction"] == operation["business_correction"]
    for rows in (_csv_rows(client), _xlsx_rows(client)):
        exported = [row for row in rows if row["operation_id"] == operation_id]
        assert len(exported) == 1
        assert exported[0]["row_type"] == "BUSINESS_CORRECTION"
        assert exported[0]["source_id"] == correction["correction_id"]
        assert exported[0]["source_operation_id"] == str(source.operation_id)
        assert json.loads(exported[0]["source_operation_ids"]) == [str(source.operation_id)]


def test_reservation_only_correction_is_item_filtered_without_physical_log(
    client, db_session, make_item,
) -> None:
    env, _request, _allocation, source = _reservation_only_prepare(
        db_session, make_item, after_observation=False,
    )
    correction = requested(db_session, env, source)
    result = approve(db_session, env[1], correction, approval_preview(db_session, env[1], correction))
    assert db_session.query(models.TransactionLog).count() == 0
    listing = client.get("/api/inventory/operations", params={"item_id": str(env[2].item_id)})
    assert listing.status_code == 200, listing.text
    operation = next(row for row in listing.json()["items"] if row["operation_id"] == result["operation_id"])
    assert operation["business_correction"]["source_operation_id"] == str(source.operation_id)
    assert operation["business_correction"]["source_operation_ids"] == [str(source.operation_id)]
    assert operation["lines"] == []
    assert [row["row_type"] for row in _csv_rows(client) if row["operation_id"] == result["operation_id"]] == [
        "BUSINESS_CORRECTION",
    ]


def test_document_correction_shows_all_valid_workflow_operations(client, db_session, make_item) -> None:
    """A picked-up shipment has both prepare and pickup operations in its correction lineage."""
    env = setup_round(db_session, make_item)
    actor, item = env[1], env[2]
    request = models.ShippingRequest(base_pf_item_id=item.item_id,
        status=models.ShippingRequestStatusEnum.PICKED_UP)
    db_session.add(request)
    db_session.flush()
    allocation = models.ShippingAllocation(request_id=request.request_id, item_id=item.item_id,
        quantity=2, department=None, status="CONSUMED")
    db_session.add(allocation)
    db_session.flush()
    prepare = inventory_operations.create_business_operation(db_session, domain="shipping", action="prepare",
        display_label="Prepared shipment", actor_name=actor.name, actor_employee_id=actor.employee_id)
    inventory_operations.record_effect(db_session, operation=prepare, effect_kind="WORKFLOW",
        subject_type="ShippingRequest", subject_id=request.request_id, role="SOURCE",
        before_state={"status": "PREPARING"}, after_state={"status": "PREPARED"})
    inventory_operations.record_effect(db_session, operation=prepare, effect_kind="ALLOCATION",
        subject_type="ShippingAllocation", subject_id=allocation.allocation_id, role="SOURCE",
        before_state={"status": "NONE"}, after_state={"status": "RESERVED"})
    db_session.flush()
    pickup = inventory_operations.create_business_operation(db_session, domain="shipping", action="pickup",
        display_label="Picked up shipment", actor_name=actor.name, actor_employee_id=actor.employee_id)
    log = models.TransactionLog(item_id=item.item_id, transaction_type=models.TransactionTypeEnum.RECEIVE,
        quantity_change=-1, quantity_before=11, quantity_after=10,
        shipping_request_id=request.request_id, shipping_phase="pickup",
        inventory_effect=[{"scope": "warehouse", "delta": -1}])
    db_session.add(inventory_operations.attach_transaction(log, pickup, models.InventoryOperationRoleEnum.PRIMARY))
    inventory_operations.record_effect(db_session, operation=pickup, effect_kind="WORKFLOW",
        subject_type="ShippingRequest", subject_id=request.request_id, role="SOURCE",
        before_state={"status": "PREPARED"}, after_state={"status": "PICKED_UP"})
    inventory_operations.record_effect(db_session, operation=pickup, effect_kind="ALLOCATION",
        subject_type="ShippingAllocation", subject_id=allocation.allocation_id, role="SOURCE",
        before_state={"status": "RESERVED"}, after_state={"status": "CONSUMED"})
    db_session.commit()
    apply(db_session, env, preview(db_session, env, confirm(db_session, env)))
    correction = stocktake.execute(db_session, actor, "correction_request", "multi-operation-request", {
        "source_type": "ShippingRequest", "source_id": str(request.request_id),
        "target_status": "PREPARING", "reason": "Dispatch did not happen",
        "evidence": "Count and both shipment phases verified"})
    result = approve(db_session, actor, correction, approval_preview(db_session, actor, correction))
    claims = {(row.source_kind, str(row.source_id)) for row in db_session.query(models.BusinessCorrectionSource)
              .filter_by(correction_id=correction["correction_id"]).all()}
    expected = sorted([str(prepare.operation_id), str(pickup.operation_id)])
    assert ("operation", str(pickup.operation_id)) in claims
    assert ("workflow_operation", str(prepare.operation_id)) in claims

    listing = client.get("/api/inventory/operations", params={"item_id": str(item.item_id)})
    assert listing.status_code == 200, listing.text
    operation = next(row for row in listing.json()["items"] if row["operation_id"] == result["operation_id"])
    assert operation["business_correction"]["source_operation_ids"] == expected
    assert operation["business_correction"]["source_operation_id"] is None
    detail = client.get(f"/api/inventory/operations/{result['operation_id']}")
    assert detail.status_code == 200, detail.text
    assert detail.json()["business_correction"] == operation["business_correction"]
    for rows in (_csv_rows(client), _xlsx_rows(client)):
        exported = next(row for row in rows if row["operation_id"] == result["operation_id"]
                        and row["row_type"] == "BUSINESS_CORRECTION")
        assert exported["source_operation_id"] in ("", None)
        assert json.loads(exported["source_operation_ids"]) == expected
