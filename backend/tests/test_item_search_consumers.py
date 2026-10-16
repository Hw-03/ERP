"""Consumer contracts preserve field scope, snapshots, and page boundaries."""
from datetime import datetime, timedelta
from uuid import UUID

import pytest

from app.models import Employee, ItemAlias, ItemSynonymGroup, ItemSynonymTerm, ShippingRequest, ShippingRequestStatusEnum, TransactionLog, TransactionTypeEnum
from app.services import stocktake_reads
from app.services.inventory_conflicts import InventoryConflict
from tests.services.test_stocktake_counts import setup_round


def dictionary(db, item, alias="Power Cell"):
    group = ItemSynonymGroup()
    db.add(group)
    db.flush()
    db.add_all([ItemSynonymTerm(group_id=group.group_id, term=term, normalized_term=term.lower())
                for term in ["Battery", "Cell"]])
    db.add(ItemAlias(item_id=item.item_id, alias=alias, normalized_alias=alias.lower().replace(" ", "")))
    db.commit()


def test_transaction_list_and_exports_apply_synonyms_only_to_item_fields(client, db_session, make_item):
    item = make_item(name="Battery")
    unrelated = make_item(name="Different")
    dictionary(db_session, item)
    db_session.add_all([TransactionLog(item_id=row.item_id, transaction_type=TransactionTypeEnum.RECEIVE,
        quantity_change=1, produced_by="Battery", created_at=datetime.utcnow()) for row in [item, unrelated]])
    db_session.commit()
    response = client.get("/api/inventory/transactions", params={"search": "Cell"})
    assert [row["item_id"] for row in response.json()] == [str(item.item_id)]
    today = (datetime.utcnow() + timedelta(hours=9)).date().isoformat()
    for extension in ["csv", "xlsx"]:
        response = client.get(f"/api/inventory/transactions/export.{extension}",
            params={"search": "Cell", "start_date": today, "end_date": today})
        assert response.status_code == 200
        if extension == "csv":
            assert "Battery" in response.content.decode("utf-8-sig")
            assert "Different" not in response.content.decode("utf-8-sig")


@pytest.mark.parametrize("search", ["Battery R123", "Cell", "Cell R123"])
def test_history_snapshot_search_matches_list_groups_and_exports(client, db_session, make_item, search):
    """Renamed item evidence participates in the same token and synonym contract."""
    item = make_item(name="NewPart")
    dictionary(db_session, item, alias="Unrelated alias")
    log = TransactionLog(item_id=item.item_id, transaction_type=TransactionTypeEnum.RECEIVE,
        quantity_change=1, reference_no="R123", created_at=datetime.utcnow(),
        item_snapshot={"item_name": "Battery", "mes_code": "OLD-001", "unit": "EA", "process_type_code": "TR"})
    db_session.add(log)
    db_session.commit()
    response = client.get("/api/inventory/transactions", params={"search": search})
    assert response.status_code == 200, response.text
    assert [row["log_id"] for row in response.json()] == [str(log.log_id)]
    response = client.get("/api/inventory/transactions/display-groups", params={"search": search})
    assert response.status_code == 200, response.text
    assert [row["log_id"] for group in response.json()["groups"] for row in group["logs"]] == [str(log.log_id)]
    today = (datetime.utcnow() + timedelta(hours=9)).date().isoformat()
    for extension in ["csv", "xlsx"]:
        response = client.get(f"/api/inventory/transactions/export.{extension}",
            params={"search": search, "start_date": today, "end_date": today})
        assert response.status_code == 200, response.text
        if extension == "csv":
            assert "Battery" in response.content.decode("utf-8-sig")
        else:
            from io import BytesIO
            from openpyxl import load_workbook
            book = load_workbook(BytesIO(response.content), read_only=True)
            try:
                assert any("Battery" in row for row in book.active.iter_rows(values_only=True))
            finally:
                book.close()


def test_shipping_alias_scope_preserves_final_custom_base_coalesce(client, db_session, make_item):
    actor = Employee(employee_code="SEARCH-SHIPPING", name="Search operator", role="worker", department="조립", is_active="true")
    db_session.add(actor)
    headers = {"X-MES-Employee-Code": actor.employee_code}
    base = make_item(name="Base", process_type_code="PF")
    final = make_item(name="Final", process_type_code="PF")
    dictionary(db_session, base)
    db_session.add(ItemAlias(item_id=final.item_id, alias="Final Cell", normalized_alias="finalcell"))
    rows = [ShippingRequest(base_pf_item_id=base.item_id, final_pf_item_id=final_id, custom_pf_name=custom,
        status=ShippingRequestStatusEnum.CANCELLED, cancelled_at=datetime.utcnow())
        for final_id, custom in [(None, None), (None, "Custom name"), (final.item_id, "Custom name")]]
    db_session.add_all(rows)
    db_session.commit()
    response = client.get("/api/shipping/history", params={"q": "Power"}, headers=headers)
    assert response.status_code == 200, response.text
    assert [row["request_id"] for row in response.json()["requests"]] == [str(rows[0].request_id)]
    response = client.get("/api/shipping/history", params={"q": "Final Cell"}, headers=headers)
    assert response.status_code == 200, response.text
    assert [row["request_id"] for row in response.json()["requests"]] == [str(rows[2].request_id)]
    page = client.get("/api/shipping/history", params={"limit": 1}, headers=headers)
    assert page.status_code == 200, page.text
    cursor = page.json()["next_cursor"]
    db_session.add(ItemAlias(item_id=base.item_id, alias="Changed", normalized_alias="changed"))
    db_session.commit()
    stale = client.get("/api/shipping/history", params={"limit": 1, "cursor": cursor}, headers=headers)
    assert stale.status_code == 409 and stale.json()["detail"]["code"] == "SEARCH_CONFIG_STALE"


def test_stocktake_snapshot_search_and_config_change_restart_cursor(db_session, make_item):
    env = setup_round(db_session, make_item, new_policy=True)
    round_id = UUID(env[3]["round_id"])
    env[2].item_name = "Current Battery"
    db_session.commit()
    dictionary(db_session, env[2])
    assert len(stocktake_reads.targets(db_session, round_id, search="Cell")["items"]) == 2
    assert len(stocktake_reads.targets(db_session, round_id, search="Counted item")["items"]) == 2
    cursor = stocktake_reads.targets(db_session, round_id, limit=1)["next_cursor"]
    db_session.add(ItemAlias(item_id=env[2].item_id, alias="Changed", normalized_alias="changed"))
    db_session.commit()
    with pytest.raises(InventoryConflict) as error:
        stocktake_reads.targets(db_session, round_id, cursor=cursor, limit=1)
    assert error.value.code == "TARGET_CURSOR_STALE"
